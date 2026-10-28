import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_rm_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw_Materials_Template"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    primary_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    instruct_font = Font(name="Segoe UI", size=9, italic=True, color="4B5563")
    data_font = Font(name="Segoe UI", size=10, color="1F2937")
    example_font = Font(name="Segoe UI", size=10, italic=True, color="047857")
    example_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid") # Light emerald

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "Material Name*",
        "Item Code*",
        "Category*",
        "Unit*",
        "HSN / SAC",
        "Opening Stock",
        "Minimum Stock",
        "Reorder Level",
        "Purchase Price (INR)",
        "Transport Charge (INR)",
        "GST %",
        "Preferred Vendor"
    ]

    hints = [
        "Required. Name of raw material",
        "Required. Unique code (e.g. RM-ALU-01)",
        "Required. Must match active Category name exactly",
        "Required. Must match active Unit symbol (e.g. PCS, KG, MTR)",
        "Optional 6-8 digit code (e.g. 76012010)",
        "Numeric >= 0 (e.g. 150.00)",
        "Numeric >= 0 (e.g. 25.00)",
        "Numeric >= 0 (e.g. 35.00)",
        "Numeric >= 0 (e.g. 245.50)",
        "Numeric >= 0 (e.g. 12.00)",
        "Numeric 0-100 (e.g. 18.00)",
        "Optional. Existing Vendor name"
    ]

    example_row = [
        "Aluminum Alloy Ingot 6063",
        "RM-ALU-01",
        "Components",
        "PCS",
        "76012010",
        150.00,
        25.00,
        35.00,
        245.50,
        12.00,
        18.00,
        "Apex Extrusions Ltd"
    ]

    # Row 1: Header
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = primary_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Row 2: Field format instructions
    hint_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    for col_idx, hint in enumerate(hints, 1):
        cell = ws.cell(row=2, column=col_idx, value=hint)
        cell.font = instruct_font
        cell.fill = hint_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[2].height = 34

    # Row 3: Example row
    for col_idx, val in enumerate(example_row, 1):
        cell = ws.cell(row=3, column=col_idx, value=val)
        cell.font = example_font
        cell.fill = example_fill
        cell.border = thin_border
        if isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    # Instructions Sheet
    ws_info = wb.create_sheet(title="Upload_Instructions")
    ws_info.views.sheetView[0].showGridLines = True
    info_title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    sec_font = Font(name="Segoe UI", size=11, bold=True, color="111827")
    bullet_font = Font(name="Segoe UI", size=10, color="374151")

    instructions = [
        ("DELUZEX ERP — BULK UPLOAD SPECIFICATIONS (RAW MATERIAL MASTER)", info_title_font),
        ("", bullet_font),
        ("1. FILE FORMAT & EXPORT GUIDELINES", sec_font),
        ("   • The Deluzex ERP Bulk Upload V1 accepts CSV (.csv) UTF-8 files.", bullet_font),
        ("   • You can enter your records in this Excel workbook, then choose 'File > Save As' and select 'CSV (Comma delimited) (*.csv)'.", bullet_font),
        ("   • Delete row 2 (Hint row) and row 3 (Sample row) before exporting, or clear their values so only real rows are uploaded.", bullet_font),
        ("", bullet_font),
        ("2. MAXIMUM CAPACITY & CONSTRAINTS", sec_font),
        ("   • Maximum data rows: 1,000 items per file.", bullet_font),
        ("   • Maximum file size: 2.00 MB.", bullet_font),
        ("", bullet_font),
        ("3. REQUIRED FIELDS (*)", sec_font),
        ("   • Material Name: Full descriptive name of the raw material.", bullet_font),
        ("   • Item Code: Unique code across active raw materials. Duplicate codes are automatically SKIPPED.", bullet_font),
        ("   • Category: Must match an existing, active Category name in Deluzex ERP (case-insensitive).", bullet_font),
        ("   • Unit: Must match an active Measurement Unit symbol (e.g. PCS, KG, MTR, LTR).", bullet_font),
        ("", bullet_font),
        ("4. STOCK & NUMERIC RULES", sec_font),
        ("   • Opening Stock, Minimum Stock, Reorder Level, Purchase Price, and Transport Charge must be non-negative numbers.", bullet_font),
        ("   • Opening stock directly initializes stock quantity on catalog creation without generating stock movement ledger entries.", bullet_font),
        ("   • GST % must be between 0.00 and 100.00 (default 18%).", bullet_font),
        ("", bullet_font),
        ("5. DUPLICATE & ERROR HANDLING", sec_font),
        ("   • Duplicate Policy is 'SKIP'. Existing items are never modified or overwritten.", bullet_font),
        ("   • If an item code already exists in the system or appears twice in the same CSV, it will be skipped and reported.", bullet_font),
        ("   • In the UI preview, you can click any cell to edit it inline or check 'Skip invalid rows' to import only valid items.", bullet_font),
    ]

    for row_idx, (text, f_style) in enumerate(instructions, 1):
        cell = ws_info.cell(row=row_idx, column=1, value=text)
        cell.font = f_style
    ws_info.column_dimensions['A'].width = 110

    # Auto-fit column widths for template sheet
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(headers[col_idx-1]), len(str(example_row[col_idx-1])), len(hints[col_idx-1]) // 2)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

    wb.save(r"D:\ERP\project-notes\bulk-upload-v1\Raw_Material_Master_Bulk_Upload_Template.xlsx")
    print("Created Raw Material Template XLSX")

def create_fp_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product_Master_Template"
    ws.views.sheetView[0].showGridLines = True

    primary_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    instruct_font = Font(name="Segoe UI", size=9, italic=True, color="4B5563")
    example_font = Font(name="Segoe UI", size=10, italic=True, color="047857")
    example_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = [
        "Product Name*",
        "Item Code / SKU*",
        "Category*",
        "Unit*",
        "HSN / SAC",
        "Opening Stock",
        "Minimum Stock",
        "Cost Price (INR)",
        "Transport Charge (INR)",
        "Dealer Selling Price (INR)",
        "Customer Selling Price (INR)",
        "GST %"
    ]

    hints = [
        "Required. Full name of finished product",
        "Required. Unique SKU code (e.g. FP-SLZ-101)",
        "Required. Must match active Category name exactly",
        "Required. Must match active Unit symbol (e.g. PCS, NOS, SET)",
        "Optional 6-8 digit code (e.g. 94054090)",
        "Numeric >= 0 (e.g. 50.00)",
        "Numeric >= 0 (e.g. 10.00)",
        "Numeric >= 0 (e.g. 850.00)",
        "Numeric >= 0 (e.g. 25.00)",
        "Numeric >= 0 (e.g. 1200.00)",
        "Numeric >= 0 (e.g. 1650.00)",
        "Numeric 0-100 (e.g. 18.00)"
    ]

    example_row = [
        "Architectural Recessed Spotlight 15W",
        "FP-SLZ-101",
        "Components",
        "PCS",
        "94054090",
        50.00,
        10.00,
        850.00,
        25.00,
        1200.00,
        1650.00,
        18.00
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = primary_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    hint_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    for col_idx, hint in enumerate(hints, 1):
        cell = ws.cell(row=2, column=col_idx, value=hint)
        cell.font = instruct_font
        cell.fill = hint_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[2].height = 34

    for col_idx, val in enumerate(example_row, 1):
        cell = ws.cell(row=3, column=col_idx, value=val)
        cell.font = example_font
        cell.fill = example_fill
        cell.border = thin_border
        if isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = '#,##0.00'
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 22

    # Instructions Sheet
    ws_info = wb.create_sheet(title="Upload_Instructions")
    ws_info.views.sheetView[0].showGridLines = True
    info_title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    sec_font = Font(name="Segoe UI", size=11, bold=True, color="111827")
    bullet_font = Font(name="Segoe UI", size=10, color="374151")

    instructions = [
        ("DELUZEX ERP — BULK UPLOAD SPECIFICATIONS (PRODUCT MASTER)", info_title_font),
        ("", bullet_font),
        ("1. FILE FORMAT & EXPORT GUIDELINES", sec_font),
        ("   • The Deluzex ERP Bulk Upload V1 accepts CSV (.csv) UTF-8 files.", bullet_font),
        ("   • You can enter your records in this Excel workbook, then choose 'File > Save As' and select 'CSV (Comma delimited) (*.csv)'.", bullet_font),
        ("   • Delete row 2 (Hint row) and row 3 (Sample row) before exporting, or clear their values so only real rows are uploaded.", bullet_font),
        ("", bullet_font),
        ("2. MAXIMUM CAPACITY & CONSTRAINTS", sec_font),
        ("   • Maximum data rows: 1,000 items per file.", bullet_font),
        ("   • Maximum file size: 2.00 MB.", bullet_font),
        ("", bullet_font),
        ("3. REQUIRED FIELDS (*)", sec_font),
        ("   • Product Name: Full descriptive SKU or catalog title.", bullet_font),
        ("   • Item Code / SKU: Unique product code across active products. Duplicate codes are automatically SKIPPED.", bullet_font),
        ("   • Category: Must match an existing, active Category name in Deluzex ERP (case-insensitive).", bullet_font),
        ("   • Unit: Must match an active Measurement Unit symbol (e.g. PCS, NOS, SET).", bullet_font),
        ("", bullet_font),
        ("4. MULTI-TIER PRICING & STOCK RULES", sec_font),
        ("   • Cost Price, Dealer Selling Price, and Customer Selling Price must be non-negative numbers.", bullet_font),
        ("   • Opening Stock directly initializes catalog stock without stock movement ledger side-effects.", bullet_font),
        ("   • GST % must be between 0.00 and 100.00 (default 18%).", bullet_font),
        ("", bullet_font),
        ("5. DUPLICATE & ERROR HANDLING", sec_font),
        ("   • Duplicate Policy is 'SKIP'. Existing product records are never updated or overwritten.", bullet_font),
        ("   • Duplicate codes inside the same batch or matching active database codes will be safely skipped.", bullet_font),
        ("   • The interactive preview allows inline editing and revalidation before finalizing submission.", bullet_font),
    ]

    for row_idx, (text, f_style) in enumerate(instructions, 1):
        cell = ws_info.cell(row=row_idx, column=1, value=text)
        cell.font = f_style
    ws_info.column_dimensions['A'].width = 110

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(headers[col_idx-1]), len(str(example_row[col_idx-1])), len(hints[col_idx-1]) // 2)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

    wb.save(r"D:\ERP\project-notes\bulk-upload-v1\Product_Master_Bulk_Upload_Template.xlsx")
    print("Created Product Master Template XLSX")

if __name__ == "__main__":
    create_rm_template()
    create_fp_template()
